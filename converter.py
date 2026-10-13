"""
Core logic: build a Nykaa InvoiceList from
  * the Nykaa PO export (PoAlpha_*.xls)      -> columns A-J
  * the sales (Tally) invoice + Master file  -> columns K-V

Handles three input flavours:
  * Excel 2003 XML ("SpreadsheetML") - what Nykaa actually exports with a .xls name
  * real binary .xls (via xlrd)
  * .xlsx (via openpyxl)

Columns are found by their header text (not fixed positions), so small layout
changes in the PO export don't break the tool.
"""
from __future__ import annotations

import io
import re
import xml.etree.ElementTree as ET
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

TEMPLATE_PATH = Path(__file__).with_name("invoice_template.xlsx")

OUTPUT_HEADERS = [
    "PO code*", "PO Line No*", "SKU code*", "Product Name*", "Vendor SKU Code*",
    "EAN Code", "HSN*", "PO Product MRP*", "PO Product Cost*", "PO Product Qty*",
    "Invoice No*", "Invoice Date*", "Invoice Qty*", "Invoice MRP*", "Invoice Cost*",
    "Discount", "Tax Rate*", "Line Tax Amount*", "CGST*", "SGST*", "IGST*", "Line Amount*",
]

# ----------------------------------------------------------------------------
# 1. Load any input file into a simple grid: {(row, col): "text"} (1-based)
# ----------------------------------------------------------------------------
_SS = "{urn:schemas-microsoft-com:office:spreadsheet}"


def _grid_from_spreadsheetml(data: bytes) -> dict:
    root = ET.fromstring(data)
    ws = root.find(f"{_SS}Worksheet")
    table = ws.find(f"{_SS}Table")
    grid, r = {}, 0
    for row in table.findall(f"{_SS}Row"):
        r = int(row.get(f"{_SS}Index", r + 1))
        c = 0
        for cell in row.findall(f"{_SS}Cell"):
            c = int(cell.get(f"{_SS}Index", c + 1))
            d = cell.find(f"{_SS}Data")
            if d is not None:
                text = "".join(d.itertext()).strip()
                if text:
                    grid[(r, c)] = text
            c += int(cell.get(f"{_SS}MergeAcross", 0))
        # rows spanning multiple lines (MergeDown) don't affect indexing of later rows
    return grid


def _grid_from_xlsx(data: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    grid = {}
    for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
        for c, v in enumerate(row, start=1):
            if v is not None and str(v).strip():
                grid[(r, c)] = str(v).strip()
    return grid


def _grid_from_xls(data: bytes) -> dict:
    import xlrd  # only needed for genuine binary .xls files

    book = xlrd.open_workbook(file_contents=data)
    sh = book.sheet_by_index(0)
    grid = {}
    for r in range(sh.nrows):
        for c in range(sh.ncols):
            v = sh.cell_value(r, c)
            if v not in ("", None):
                if isinstance(v, float) and v.is_integer():
                    v = int(v)
                grid[(r + 1, c + 1)] = str(v).strip()
    return grid


def load_grid(data: bytes) -> dict:
    head = data[:512].lstrip()
    if head.startswith(b"<?xml") or b"urn:schemas-microsoft-com:office:spreadsheet" in data[:4000]:
        return _grid_from_spreadsheetml(data)
    if data[:2] == b"PK":
        return _grid_from_xlsx(data)
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return _grid_from_xls(data)
    raise ValueError("Unrecognised file format. Please upload the PO Excel file (.xls / .xlsx).")


# ----------------------------------------------------------------------------
# 2. Parse the PO
# ----------------------------------------------------------------------------
@dataclass
class POLine:
    line_no: int
    sku_code: str
    product_name: str
    vendor_sku: str
    ean: str
    hsn: str
    mrp: float
    cost: float
    qty: int
    tax_rate: float      # total GST % (CGST+SGST+IGST)
    is_igst: bool        # inter-state supply -> tax goes to IGST column


@dataclass
class PurchaseOrder:
    po_number: str
    po_date: str = ""
    supplier: str = ""
    lines: list[POLine] = field(default_factory=list)


def _num(s, default=0.0) -> float:
    if s is None:
        return default
    s = str(s).replace(",", "").replace("%", "").replace("₹", "").strip()
    try:
        return float(s)
    except ValueError:
        return default


def _clean_name(s: str) -> str:
    # Nykaa appends "\nColour : \nSize :" to every product name
    s = re.split(r"\n\s*Colour\s*:", s or "")[0]
    return " ".join(s.split())


def _find(grid, pattern, after_row=0):
    rx = re.compile(pattern, re.I)
    for (r, c), v in sorted(grid.items()):
        if r > after_row and rx.fullmatch(v.strip()):
            return r, c
    return None


def _value_right_of(grid, r, c):
    for cc in range(c + 1, c + 30):
        if (r, cc) in grid:
            return grid[(r, cc)]
    return ""


def parse_po(data: bytes) -> PurchaseOrder:
    grid = load_grid(data)

    pos = _find(grid, r"PO\s*No\.?")
    if not pos:
        raise ValueError("Could not find 'PO No' in the file - is this a Nykaa PO export?")
    po = PurchaseOrder(po_number=_value_right_of(grid, *pos))
    if (d := _find(grid, r"PO\s*Date")):
        po.po_date = _value_right_of(grid, *d)
    if (s := _find(grid, r"Sold By")):
        po.supplier = grid.get((s[0] + 1, s[1]), "") or grid.get((s[0] + 8, s[1]), "")

    hdr = _find(grid, r"Item\s*Code")
    if not hdr:
        raise ValueError("Could not find the item table (header 'Item Code').")
    hr = hdr[0]
    cols = {v.strip().lower(): c for (r, c), v in grid.items() if r == hr}

    def col(*names):
        for n in names:
            if n.lower() in cols:
                return cols[n.lower()]
        return None

    c_no, c_sku, c_ean = col("#", "sr no", "s.no"), col("item code"), col("ean")
    c_vsku, c_name, c_hsn = col("vendor sku", "vendor sku code"), col("sku name", "product name"), col("hsn")
    c_qty, c_mrp, c_cost = col("qty", "quantity"), col("mrp"), col("unit price", "cost")
    c_cgst, c_sgst, c_igst = col("cgst"), col("sgst/ugst", "sgst"), col("igst")
    missing = [n for n, c in [("#", c_no), ("Item Code", c_sku), ("Qty", c_qty), ("MRP", c_mrp),
                              ("Unit Price", c_cost)] if c is None]
    if missing:
        raise ValueError(f"PO item table is missing columns: {', '.join(missing)}")

    r = hr + 1
    max_r = max(k[0] for k in grid)
    started = False
    while r <= max_r:
        no = grid.get((r, c_no), "")
        if re.fullmatch(r"\d+", no):
            started = True
            g = lambda c: grid.get((r, c), "") if c else ""
            cg, sg, ig = _num(g(c_cgst)), _num(g(c_sgst)), _num(g(c_igst))
            po.lines.append(POLine(
                line_no=int(no),
                sku_code=g(c_sku),
                product_name=_clean_name(g(c_name)),
                vendor_sku=g(c_vsku),
                ean=g(c_ean),
                hsn=g(c_hsn),
                mrp=round(_num(g(c_mrp)), 2),
                cost=round(_num(g(c_cost)), 4),
                qty=int(_num(g(c_qty))),
                tax_rate=round(cg + sg + ig, 2),
                is_igst=ig > 0,
            ))
        elif started:
            break  # first non-numeric row after items = totals row
        r += 1

    if not po.lines:
        raise ValueError("No item lines found in the PO.")
    return po


# ----------------------------------------------------------------------------
# 3. Sales (Tally) invoice and Master file  -> columns K..V
# ----------------------------------------------------------------------------
def norm_ean(v) -> str:
    """'8904055020667', 8904055020667, 8904055020667.0 -> '8904055020667'"""
    if v is None:
        return ""
    s = str(v).strip()
    s = re.sub(r"\.0+$", "", s)
    if re.fullmatch(r"\d+(\.\d+)?[eE]\+\d+", s):  # 8.904055020667E+12
        s = str(int(float(s)))
    return s


@dataclass
class InvoiceLine:
    row: int             # row number in the sales invoice (for messages)
    name: str
    ean: str
    qty: float
    rate: float          # GST-inclusive rate before discount
    disc_pct: float
    gst_pct: float


@dataclass
class SalesInvoice:
    invoice_no: str
    invoice_date: str     # DD/MM/YYYY
    po_number: str
    is_igst: bool
    lines: list[InvoiceLine] = field(default_factory=list)


def parse_sales_invoice(data: bytes) -> SalesInvoice:
    grid = load_grid(data)

    def search(pattern):
        rx = re.compile(pattern, re.I)
        for (_r, _c), v in sorted(grid.items()):
            m = rx.search(v)
            if m:
                return m
        return None

    m_no = search(r"Invoice\s*No\.?\s*:\s*([^\s]+)")
    m_dt = search(r"\bDate\s*:\s*(\d{1,2})[./-](\d{1,2})[./-](\d{4})")
    m_po = search(r"PO\s*No\.?\s*:\s*([^\s]+)")
    if not m_no:
        raise ValueError("Could not find 'Invoice No. :' in the sales invoice.")
    if not m_dt:
        raise ValueError("Could not find 'Date :' (DD.MM.YYYY) in the sales invoice.")
    inv_date = f"{int(m_dt.group(1)):02d}/{int(m_dt.group(2)):02d}/{m_dt.group(3)}"

    hdr = _find(grid, r"EAN\s*Code")
    if not hdr:
        raise ValueError("Could not find the item table (header 'EAN Code') in the sales invoice.")
    hr = hdr[0]
    cols = {v.strip().lower(): c for (r, c), v in grid.items() if r == hr}
    need = {"ean code": None, "qty nos": None, "rate": None, "disc %": None, "gst%": None}
    for k in need:
        need[k] = cols.get(k)
    missing = [k for k, v in need.items() if v is None]
    if missing:
        raise ValueError(f"Sales invoice item table is missing columns: {', '.join(missing)}")
    c_name = cols.get("particular", 1)
    is_igst = "igst amount" in cols

    inv = SalesInvoice(invoice_no=m_no.group(1).strip(), invoice_date=inv_date,
                       po_number=m_po.group(1).strip() if m_po else "", is_igst=is_igst)
    max_r = max(k[0] for k in grid)
    for r in range(hr + 1, max_r + 1):
        name = grid.get((r, c_name), "")
        if re.match(r"\s*(gross\s*)?total", name, re.I):
            break
        ean = norm_ean(grid.get((r, need["ean code"])))
        qty = _num(grid.get((r, need["qty nos"])), None)
        if not ean and qty is None:
            continue  # blank spacer row
        inv.lines.append(InvoiceLine(
            row=r, name=" ".join(name.split()), ean=ean, qty=qty or 0,
            rate=_num(grid.get((r, need["rate"]))),
            disc_pct=_num(grid.get((r, need["disc %"]))),
            gst_pct=_num(grid.get((r, need["gst%"]))),
        ))
    if not inv.lines:
        raise ValueError("No item lines found in the sales invoice.")
    return inv


def load_master(files: list[tuple[str, bytes]]) -> tuple[dict, list[str]]:
    """Master file(s) with columns 'EAN' and 'MRP' -> ({ean: mrp}, messages)."""
    import pandas as pd

    mrp, msgs = {}, []
    for name, data in files:
        try:
            df = pd.read_excel(io.BytesIO(data), dtype={"EAN": str})
        except Exception as e:
            msgs.append(f"{name}: could not read ({e})")
            continue
        df.columns = [str(c).strip() for c in df.columns]
        upper = {c.upper(): c for c in df.columns}
        if "EAN" not in upper or "MRP" not in upper:
            msgs.append(f"{name}: skipped - columns 'EAN' and 'MRP' not found")
            continue
        n = 0
        for e, m in zip(df[upper["EAN"]], df[upper["MRP"]]):
            e = norm_ean(e)
            if e and e.lower() != "nan" and e not in mrp and pd.notna(m):
                try:
                    mrp[e] = float(m)
                    n += 1
                except (TypeError, ValueError):
                    pass
        msgs.append(f"{name}: {n} EANs loaded")
    return mrp, msgs


def invoice_values(il: InvoiceLine, is_igst: bool, mrp) -> dict:
    """K..V values - same formulas as the Nykaa Invoice Filler (main.py)."""
    cost = il.rate * (1 - il.disc_pct / 100) / (1 + il.gst_pct / 100)
    tax = il.qty * cost * il.gst_pct / 100
    return dict(
        invoice_qty=int(il.qty) if float(il.qty).is_integer() else il.qty,
        invoice_mrp=mrp if mrp is not None else "",
        invoice_cost=cost,
        discount=il.disc_pct / 100,
        tax_rate=int(il.gst_pct) if float(il.gst_pct).is_integer() else il.gst_pct,
        line_tax=tax,
        cgst=0 if is_igst else tax / 2,
        sgst=0 if is_igst else tax / 2,
        igst=tax if is_igst else 0,
        line_amount=il.qty * cost + tax,
    )


def auto_match(po: PurchaseOrder, inv: SalesInvoice) -> tuple[dict, list[int], list[int]]:
    """
    Match invoice lines to PO lines by EAN.
    Returns ({invoice_idx: po_idx}, unmatched_invoice_idx, unmatched_po_idx)
    """
    po_by_ean = {}
    for i, ln in enumerate(po.lines):
        po_by_ean.setdefault(norm_ean(ln.ean), i)
    pairs, used = {}, set()
    for j, il in enumerate(inv.lines):
        i = po_by_ean.get(il.ean)
        if i is not None and i not in used:
            pairs[j] = i
            used.add(i)
    un_inv = [j for j in range(len(inv.lines)) if j not in pairs]
    un_po = [i for i in range(len(po.lines)) if i not in used]
    return pairs, un_inv, un_po


def suggest(il: InvoiceLine, po: PurchaseOrder, candidates: list[int]) -> tuple[int | None, float]:
    """Best PO line for an unmatched invoice line, by product name and EAN similarity."""
    from difflib import SequenceMatcher

    def simplify(s):
        s = s.lower().replace("&", " and ").replace("aloevera", "aloe vera").replace("alovera", "aloe vera")
        return " ".join(re.findall(r"[a-z0-9]+", s))

    best, score = None, 0.0
    for i in candidates:
        p = po.lines[i]
        name_score = SequenceMatcher(None, simplify(il.name), simplify(p.product_name)[: len(simplify(il.name)) + 15]).ratio()
        ean_score = SequenceMatcher(None, il.ean, norm_ean(p.ean)).ratio() if il.ean and p.ean else 0
        s = max(name_score, ean_score if ean_score >= 0.9 else 0)
        if s > score:
            best, score = i, s
    return best, score


# ----------------------------------------------------------------------------
# 4. Build the InvoiceList workbook
# ----------------------------------------------------------------------------
# Column order A..J of the InvoiceList sheet
PO_FIELDS = ["po_code", "line_no", "sku_code", "product_name", "vendor_sku",
             "ean", "hsn", "mrp", "cost", "qty"]
INV_FIELDS = ["invoice_no", "invoice_date", "invoice_qty", "invoice_mrp", "invoice_cost",
              "discount", "tax_rate", "line_tax", "cgst", "sgst", "igst", "line_amount"]
_DEC = "0.00;-0.00;0"
INV_FORMATS = {"invoice_mrp": _DEC, "invoice_cost": _DEC, "discount": "0%", "tax_rate": "General",
               "line_tax": _DEC, "cgst": _DEC, "sgst": _DEC, "igst": _DEC, "line_amount": _DEC}


def build_invoice_list(rows: list[dict], template_path: Path = TEMPLATE_PATH) -> bytes:
    """
    rows: dicts with the A-J keys (PO_FIELDS) and optionally the K-V keys (INV_FIELDS).
    Missing K-V keys are left empty. Returns the .xlsx file as bytes.
    """
    wb = openpyxl.load_workbook(template_path)
    ws = wb["InvoiceList"]

    body_style = {c: copy(ws.cell(2, c)._style) for c in range(1, 23)}
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row)
    ws.data_validations.dataValidation = []

    center = Alignment(horizontal="center", vertical="center")
    keys = PO_FIELDS + INV_FIELDS
    for i, row in enumerate(rows, start=2):
        for c, k in enumerate(keys, start=1):
            v = row.get(k)
            if k == "hsn" and str(v).isdigit():
                v = int(v)  # HSN as a number, like Nykaa's sample
            cell = ws.cell(i, c, v)
            cell._style = copy(body_style[c])
            cell.alignment = center
            if k in INV_FORMATS and v not in (None, ""):
                cell.number_format = INV_FORMATS[k]

    wb.active = wb.sheetnames.index("InvoiceList")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def po_to_rows(po: PurchaseOrder, vendor_sku_fallback: str = "blank") -> list[dict]:
    """One dict per PO line with the A-J fields."""
    out = []
    for ln in po.lines:
        vsku = ln.vendor_sku
        if not vsku and vendor_sku_fallback == "sku":
            vsku = ln.sku_code
        elif not vsku and vendor_sku_fallback == "ean":
            vsku = ln.ean
        out.append(dict(
            po_code=po.po_number, line_no=ln.line_no, sku_code=ln.sku_code,
            product_name=ln.product_name, vendor_sku=vsku, ean=ln.ean, hsn=ln.hsn,
            mrp=ln.mrp, cost=ln.cost, qty=ln.qty,
        ))
    return out


def output_filename(po_number: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9_-]", "", po_number) or "PO"
    return f"InvoiceList_{safe}.xlsx"
